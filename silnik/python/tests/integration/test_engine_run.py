from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from zestaw_wejsciowy import zestaw_wejsciowy

from investment_tax_engine.app.engine import InputBundle, InvestmentTaxEngine
from investment_tax_engine.models.core import EngineConfig, FundingCostEvent, TaxFilingRequest, UserOverrides


# Testy uzywaja prawdziwych wyciagow, gdy sa na dysku, a na czystym klonie
# zestawu syntetycznego. Wczesniej caly plik byl pomijany bez tych plikow, wiec
# osiem testow - w tym jedyny pelny przebieg PIT - nie wykonywalo sie nigdzie
# poza komputerem wlasciciela.
ZESTAW = zestaw_wejsciowy()
FIXTURE_ROOT = ZESTAW.root
DEPO_FIXTURE = ZESTAW.depo_json


def test_engine_run_on_fixture_bundle_exports_authoritative_outputs(tmp_path):
    silnik = InvestmentTaxEngine(
        EngineConfig(
            run_mode="SAFE",
            nbp_allow_api_fallback=False,
        )
    )
    bundle = InputBundle(
        api_json_path=ZESTAW.api_json,
        trades_v1_path=ZESTAW.trades_v1,
        trades_legacy_path=ZESTAW.trades_legacy,
        tradernet_table_path=ZESTAW.tradernet_table,
        broker_json_path=ZESTAW.broker_json,
        broker_xml_path=None,
        depo_json_path=DEPO_FIXTURE,
        # Bez tax_year raport obejmuje wszystkie lata, wiec komplet archiwow NBP
        # jest warunkiem koniecznym - inaczej bramka kursowa slusznie blokuje pakiet.
        nbp_csv_paths=list(ZESTAW.nbp_csv),
        output_dir=tmp_path,
    )

    result = silnik.run(bundle)

    assert result.merge_result.ledger.trades_by_id
    assert result.merge_result.ledger.events_by_id
    assert result.canonical_dataset.trades
    assert "defensible" in result.scenario_results
    assert result.primary_scenario == "aggressive_user"
    assert result.plan_used == "aggressive_user"
    # Fikstura zawiera dochod z waloru bez ustalonego kraju; PIT/ZG wymaga decyzji.
    assert result.status == "CALCULATION_BLOCKED"
    assert result.filing_ready is False
    bez_kraju = next(issue for issue in result.quality_report.blocking_issues if issue.code == "PIT_ZG_COUNTRY_UNKNOWN")
    # Komunikat wskazuje transakcje do poprawienia, a nie tylko kwote wiersza "XX".
    assert bez_kraju.details["symbols"]
    assert all(symbol in bez_kraju.message for symbol in bez_kraju.details["symbols"])
    assert result.quality_report is not None
    assert result.performance_profile
    audit_hash_before_profile_change = result.audit_hash
    result.performance_profile["test_probe"] = {"duration_ms": 1}
    assert result.audit_hash == audit_hash_before_profile_change
    assert "performance_profile" not in result.annual_summary

    workbook_path = next(path for path in result.exported_files if path.suffix == ".xlsx")
    workbook = openpyxl.load_workbook(workbook_path)

    assert "Pipeline_Status" in workbook.sheetnames
    assert "Quality_Gates" in workbook.sheetnames
    assert "Canonical_Trades" in workbook.sheetnames
    assert "Canonical_Events" in workbook.sheetnames
    assert "FX_Coverage" in workbook.sheetnames
    assert "DePo_Reconciliation" in workbook.sheetnames
    assert "FIFO_Realized" in workbook.sheetnames
    assert "FIFO_Other_Years" in workbook.sheetnames
    assert "PIT38_Summary" in workbook.sheetnames
    assert "Annual_PIT38_Summary" in workbook.sheetnames
    assert "Scenarios" in workbook.sheetnames
    assert "Private_Cash_FX" in workbook.sheetnames
    assert "Result_Health_Check" in workbook.sheetnames
    assert "Tax_Advisor_Brief" in workbook.sheetnames


def test_engine_export_perf_mode_skips_workbook_without_changing_audit_hash(tmp_path, monkeypatch):
    full_dir = tmp_path / "full"
    minimal_dir = tmp_path / "minimal"
    full_dir.mkdir()
    minimal_dir.mkdir()
    bundle_kwargs = {
        "api_json_path": ZESTAW.api_json,
        "trades_v1_path": ZESTAW.trades_v1,
        "trades_legacy_path": ZESTAW.trades_legacy,
        "tradernet_table_path": ZESTAW.tradernet_table,
        "broker_json_path": ZESTAW.broker_json,
        "broker_xml_path": None,
        "depo_json_path": DEPO_FIXTURE,
        "nbp_csv_paths": [ZESTAW.nbp_csv[0]],
    }

    monkeypatch.delenv("INVEST_TAX_EXPORT_MODE", raising=False)
    full_result = InvestmentTaxEngine(
        EngineConfig(
            run_mode="SAFE",
            nbp_allow_api_fallback=False,
        )
    ).run(InputBundle(**bundle_kwargs, output_dir=full_dir))

    monkeypatch.setenv("INVEST_TAX_EXPORT_MODE", "minimal")
    minimal_result = InvestmentTaxEngine(
        EngineConfig(
            run_mode="SAFE",
            nbp_allow_api_fallback=False,
        )
    ).run(InputBundle(**bundle_kwargs, output_dir=minimal_dir))

    assert full_result.audit_hash == minimal_result.audit_hash
    assert full_result.annual_summary == minimal_result.annual_summary
    assert any(path.suffix == ".xlsx" for path in full_result.exported_files)
    assert not any(path.suffix == ".xlsx" for path in minimal_result.exported_files)
    assert any(
        stage["stage"] == "exports.workbook" and stage.get("skipped") is True
        for stage in minimal_result.performance_profile["stages"]
    )
    assert any(
        stage["stage"] == "exports.audit_json"
        for stage in minimal_result.performance_profile["stages"]
    )


def test_engine_run_for_2025_does_not_require_2026_nbp_rates(tmp_path):
    config = EngineConfig(
        run_mode="SAFE",
        nbp_allow_api_fallback=False,
    )
    config.tax_year = 2025
    silnik = InvestmentTaxEngine(config)
    bundle = InputBundle(
        api_json_path=ZESTAW.api_json,
        trades_v1_path=ZESTAW.trades_v1,
        trades_legacy_path=ZESTAW.trades_legacy,
        tradernet_table_path=ZESTAW.tradernet_table,
        broker_json_path=ZESTAW.broker_json,
        broker_xml_path=None,
        depo_json_path=DEPO_FIXTURE,
        nbp_csv_paths=[ZESTAW.nbp_csv[0]],
        output_dir=tmp_path,
    )

    result = silnik.run(bundle)

    tax_fx_issues = [issue for issue in result.merge_result.ledger.issues if issue.stage == "TAX_FX"]

    assert result.annual_summary["tax_year"] == "2025"
    assert not any(issue.code == "NBP_RATE_NOT_FOUND" for issue in tax_fx_issues)


def test_engine_run_for_2026_reports_fx_coverage_gap_without_local_2026_archive(tmp_path):
    config = EngineConfig(
        run_mode="SAFE",
        nbp_allow_api_fallback=False,
    )
    config.tax_year = 2026
    silnik = InvestmentTaxEngine(config)
    bundle = InputBundle(
        api_json_path=ZESTAW.api_json,
        trades_v1_path=ZESTAW.trades_v1,
        trades_legacy_path=ZESTAW.trades_legacy,
        tradernet_table_path=ZESTAW.tradernet_table,
        broker_json_path=ZESTAW.broker_json,
        broker_xml_path=None,
        depo_json_path=DEPO_FIXTURE,
        nbp_csv_paths=[ZESTAW.nbp_csv[0]],
        output_dir=tmp_path,
    )

    result = silnik.run(bundle)

    assert result.annual_summary["tax_year"] == "2026"
    assert result.filing_ready is False
    assert result.quality_report is not None
    assert any(gap.currency == "USD" for gap in result.fx_coverage_gaps)
    # Brak kursu NBP dla transakcji 2026 to bramka krytyczna: pakiet nie moze
    # zostac uznany za gotowy do zlozenia, mimo ze wynik jest policzony.
    assert result.status == "CALCULATION_BLOCKED"
    blocking_codes = {issue.code for issue in result.quality_report.blocking_issues}
    assert "NBP_RATE_NOT_FOUND" in blocking_codes
    assert "NBP_COVERAGE_GAP" in blocking_codes


def test_engine_run_can_generate_tax_filing_package_and_allocate_funding_fee(tmp_path):
    config = EngineConfig(
        run_mode="SAFE",
        nbp_allow_api_fallback=False,
        tax_year=2025,
        tax_filing_request=TaxFilingRequest(package_scope="full", filing_mode="CORRECTION"),
    )
    silnik = InvestmentTaxEngine(config)
    bundle = InputBundle(
        api_json_path=ZESTAW.api_json,
        trades_v1_path=ZESTAW.trades_v1,
        trades_legacy_path=ZESTAW.trades_legacy,
        tradernet_table_path=ZESTAW.tradernet_table,
        broker_json_path=ZESTAW.broker_json,
        broker_xml_path=None,
        depo_json_path=DEPO_FIXTURE,
        nbp_csv_paths=list(ZESTAW.nbp_csv),
        output_dir=tmp_path,
        user_overrides=UserOverrides(
            funding_cost_events=[
                FundingCostEvent(
                    funding_event_id="BANK-FEE-1",
                    amount=Decimal("25.00"),
                    currency="PLN",
                    date=pd.Timestamp("2025-01-09"),
                    linked_deposit_id="DEP-2025-01-09",
                    source="USER_OVERRIDE",
                    evidence_note="Bank transfer fee for investment funding",
                    deposit_amount=Decimal("10000.00"),
                )
            ]
        ),
    )

    result = silnik.run(bundle)

    assert result.tax_filing_package is not None
    assert result.tax_filing_package.draft.form_type == "PIT-38"
    assert result.tax_filing_package.draft.filing_mode == "CORRECTION"
    assert result.tax_filing_package.draft.package_scope == "full"
    assert result.tax_filing_package.justification is not None
    assert not result.funding_fee_allocations
    assert not any(item.kind == "BANK_FUNDING_FEE" for item in result.cost_items)
    assert any(issue.code == "FUNDING_FEE_DEPOSIT_UNCONFIRMED" for issue in result.actionable_issues)
    projections = result.tax_filing_package.draft.scenario_projections
    assert set(projections.keys()) == {"conservative", "defensible", "aggressive_user"}
    assert result.tax_filing_package.draft.form_fields == projections["aggressive_user"].form_fields
    aggressive_fields = {field.position: field for field in projections["aggressive_user"].form_fields}
    assert aggressive_fields["22"].value == Decimal(str(result.scenario_results["aggressive_user"].total_revenue_pln))
    assert aggressive_fields["23"].value == Decimal(str(result.scenario_results["aggressive_user"].total_cost_pln))
    assert aggressive_fields["28"].value == Decimal(str(result.scenario_results["aggressive_user"].gross_result_pln))
    assert aggressive_fields["31"].value == Decimal(str(projections["aggressive_user"].rounded_base_pln))
    assert aggressive_fields["33"].value == Decimal(str(projections["aggressive_user"].rounded_tax_from_base_pln))
    assert aggressive_fields["47"].value == Decimal(str(projections["aggressive_user"].foreign_dividend_tax_pln))
    assert aggressive_fields["48"].value == Decimal(str(projections["aggressive_user"].foreign_tax_credit_pln))

    workbook_path = next(path for path in result.exported_files if path.suffix == ".xlsx")
    workbook = openpyxl.load_workbook(workbook_path)

    assert "PIT38_Draft" in workbook.sheetnames
    assert "PIT_Form_Fields" in workbook.sheetnames
    assert "PIT38_Form_Projections" in workbook.sheetnames
    assert "Calculation_Report" in workbook.sheetnames
    assert "Justification_Memo" in workbook.sheetnames
    assert "Aggressive_Cost_Defense" in workbook.sheetnames
    assert "Funding_Fee_Allocations" in workbook.sheetnames
    defense_sheet = workbook["Aggressive_Cost_Defense"]
    header = [cell.value for cell in defense_sheet[1]]
    assert "legal_basis" in header
    assert "tax_argument_pl" in header
    assert "defense_status" in header
    assert "missing_evidence" in header
    assert "linked_trade_date" in header
    assert "amount_reconciliation" in header
    assert defense_sheet.max_row > 1
    assert "Defense_Readiness" in workbook.sheetnames
    assert "PIT_Submission_Checklist" in workbook.sheetnames
    assert "PIT_Case_File" in workbook.sheetnames
    assert "Aggressive_Memorandum" in workbook.sheetnames
    assert "Source_Manifest" in workbook.sheetnames
    assert "Import_Intelligence" in workbook.sheetnames
    assert "No_Overpay_Audit" in workbook.sheetnames
    assert "Defense_Case_File" in workbook.sheetnames
    assert "Coverage_Matrix" in workbook.sheetnames
    assert "No_Overpay_Audit_v2" in workbook.sheetnames
    assert "Defense_Case_File_v2" in workbook.sheetnames
    assert "Broker_File_Action_Queue" in workbook.sheetnames
    assert "Legal_Basis" in workbook.sheetnames
    assert "Result_Health_Check" in workbook.sheetnames
    assert "Tax_Advisor_Brief" in workbook.sheetnames
    assert "Active Sources" in workbook.sheetnames
    assert "Candidate Sources" in workbook.sheetnames
    assert "Source Trust" in workbook.sheetnames
    assert "No Overpay Guard" in workbook.sheetnames
    assert "Defense Vault" in workbook.sheetnames
    assert "Advisor Questions" in workbook.sheetnames
    assert "NBP Coverage" in workbook.sheetnames
    assert "Result Health" in workbook.sheetnames
    assert "Audit Metadata" in workbook.sheetnames
    assert any(path.name == "memorandum_aggressive_user.md" for path in result.exported_files)
    assert any(path.name == "tax_advisor_brief.md" for path in result.exported_files)
    assert any(path.name == "tax_advisor_brief.html" for path in result.exported_files)
    assert any(path.name == "tax_advisor_brief.json" for path in result.exported_files)
    assert any(path.name == "advisor_review_pack.json" for path in result.exported_files)
    assert any(path.name == "evidence_checklist.json" for path in result.exported_files)
    assert any(path.name == "evidence_checklist.xlsx" for path in result.exported_files)
    assert "tax_advisor_brief" in result.tax_filing_package.audit_appendix
    assert "result_health_check" in result.tax_filing_package.audit_appendix
    assert "advisor_review_pack" in result.tax_filing_package.audit_appendix
    assert "source_trust_summary" in result.tax_filing_package.audit_appendix


def test_real_storage_baseline_result_health_is_not_suspicious_zero(tmp_path):
    config = EngineConfig(
        run_mode="SAFE",
        nbp_allow_api_fallback=False,
        tax_year=2025,
        tax_filing_request=TaxFilingRequest(package_scope="full", filing_mode="ORIGINAL"),
    )
    silnik = InvestmentTaxEngine(config)
    bundle = InputBundle(
        api_json_path=ZESTAW.api_json,
        trades_v1_path=ZESTAW.trades_v1,
        trades_legacy_path=ZESTAW.trades_legacy,
        tradernet_table_path=ZESTAW.tradernet_table,
        broker_json_path=ZESTAW.broker_json,
        broker_xml_path=None,
        depo_json_path=DEPO_FIXTURE,
        nbp_csv_paths=list(ZESTAW.nbp_csv),
        output_dir=tmp_path,
    )

    result = silnik.run(bundle)

    assert result.tax_filing_package is not None
    health = result.tax_filing_package.audit_appendix["result_health_check"]
    assert health["activeTaxSourceLabels"], health
    assert int(health["taxHistoryRowCount"]) > 0, health
    assert int(health["sellRowCount"]) > 0, health
    assert Decimal(str(health["revenuePln"])) > 0, health
    assert Decimal(str(health["costPln"])) > 0, health
    assert health["status"] in {"ok", "needs_review"}, health
    assert not any(
        reason
        for reason in health.get("reasons", [])
        if "Brak pliku pomocniczego" in str(reason)
        or "Broker XML" in str(reason)
        or "depozytariusz" in str(reason).lower()
        or "PDF" in str(reason)
    ), health


def test_engine_run_splits_info_issues_and_projects_authoritative_history_rows(tmp_path):
    config = EngineConfig(
        run_mode="SAFE",
        nbp_allow_api_fallback=False,
        tax_year=2025,
    )
    silnik = InvestmentTaxEngine(config)
    bundle = InputBundle(
        api_json_path=ZESTAW.api_json,
        trades_v1_path=ZESTAW.trades_v1,
        trades_legacy_path=ZESTAW.trades_legacy,
        tradernet_table_path=ZESTAW.tradernet_table,
        broker_json_path=ZESTAW.broker_json,
        broker_xml_path=None,
        depo_json_path=DEPO_FIXTURE,
        nbp_csv_paths=list(ZESTAW.nbp_csv),
        output_dir=tmp_path,
        user_overrides=UserOverrides(
            funding_cost_events=[
                FundingCostEvent(
                    funding_event_id="BANK-FEE-HISTORY-1",
                    amount=Decimal("36.85"),
                    currency="USD",
                    date=pd.Timestamp("2025-01-21"),
                    linked_deposit_id="82160362",
                    source="USER_OVERRIDE",
                    evidence_note="Domyślna alokacja: proportional_first_batch",
                    deposit_amount=Decimal("10527.68"),
                )
            ]
        ),
    )

    result = silnik.run(bundle)

    assert result.informational_issues
    assert not any(issue.code == "BROKER_XML_NOT_SUPPLIED" for issue in result.actionable_issues)
    assert any(issue.code == "BROKER_XML_NOT_SUPPLIED" for issue in result.informational_issues)
    assert result.transaction_history_rows
    assert any(row.row_kind == "BONUS_CONTEST_SHARE" for row in result.transaction_history_rows)
    assert not any(
        row.row_kind == "TRADE"
        and row.transaction_id
        and str(row.transaction_id).startswith("DEPO-AWARD-")
        for row in result.transaction_history_rows
    )
    allocated_cost_rows = [row for row in result.transaction_history_rows if row.row_kind == "ALLOCATED_COST"]
    if ZESTAW.prywatny:
        # Prawdziwe wyciagi zawieraja wplate 82160362 - oplata ma potwierdzona
        # wplate, wiec rozklada sie na zakupy sfinansowane ta wplata.
        assert result.funding_fee_allocations
        assert allocated_cost_rows
        assert all(row.parent_row_id for row in allocated_cost_rows)
        assert any(row.comment and "proportional_first_batch" in row.comment for row in allocated_cost_rows)
        assert not any(issue.code == "FUNDING_FEE_DEPOSIT_UNCONFIRMED" for issue in result.actionable_issues)
    else:
        # Zestaw syntetyczny nie ma tej wplaty - bez potwierdzenia oplata nie wchodzi (fail-closed).
        assert allocated_cost_rows == []
        assert any(issue.code == "FUNDING_FEE_DEPOSIT_UNCONFIRMED" for issue in result.actionable_issues)


def test_engine_run_for_2026_uses_local_updated_csv_archives_without_coverage_gap(tmp_path):
    config = EngineConfig(
        run_mode="SAFE",
        nbp_allow_api_fallback=False,
        tax_year=2026,
    )
    silnik = InvestmentTaxEngine(config)
    bundle = InputBundle(
        api_json_path=ZESTAW.api_json,
        trades_v1_path=ZESTAW.trades_v1,
        trades_legacy_path=ZESTAW.trades_legacy,
        tradernet_table_path=ZESTAW.tradernet_table,
        broker_json_path=ZESTAW.broker_json,
        broker_xml_path=None,
        depo_json_path=DEPO_FIXTURE,
        nbp_csv_paths=list(ZESTAW.nbp_csv),
        output_dir=tmp_path,
    )

    result = silnik.run(bundle)

    assert not any(gap.currency == "USD" for gap in result.fx_coverage_gaps)


def test_fifo_sheet_of_the_filing_year_reconciles_with_the_form(tmp_path):
    """Suma arkusza FIFO musi zgadzac sie z pozycjami 22 i 23 PIT-38.

    Eksport dostawal `result.fifo_rows` bez filtra roku, choc formularz liczy
    wylacznie rok rozliczany. Kontrolujacy, ktory zsumowal arkusz pakietu
    dowodowego, dostawal inna kwote niz deklaracja i nie mial jak jej wyjasnic.
    """
    rok = 2025
    silnik = InvestmentTaxEngine(
        EngineConfig(
            run_mode="SAFE",
            tax_year=rok,
            nbp_allow_api_fallback=False,
        )
    )
    bundle = InputBundle(
        api_json_path=ZESTAW.api_json,
        trades_v1_path=ZESTAW.trades_v1,
        trades_legacy_path=ZESTAW.trades_legacy,
        tradernet_table_path=ZESTAW.tradernet_table,
        broker_json_path=ZESTAW.broker_json,
        broker_xml_path=None,
        depo_json_path=DEPO_FIXTURE,
        nbp_csv_paths=list(ZESTAW.nbp_csv),
        output_dir=tmp_path,
    )

    result = silnik.run(bundle)

    workbook_path = next(path for path in result.exported_files if path.suffix == ".xlsx")
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    arkusz = workbook["FIFO_Realized"]
    naglowki = [cell.value for cell in next(arkusz.iter_rows(max_row=1))]

    def kolumna(nazwa: str) -> list:
        indeks = naglowki.index(nazwa)
        return [row[indeks] for row in arkusz.iter_rows(min_row=2, values_only=True)]

    # Pusty arkusz spelnilby kazda asercje ponizej, wiec najpierw dowod, ze cos zawiera.
    assert arkusz.max_row > 1, "arkusz FIFO rozliczanego roku nie moze byc pusty"
    assert set(kolumna("tax_year")) <= {rok}, "arkusz rozliczanego roku nie moze miec innych lat"

    przychod = sum(Decimal(str(value)) for value in kolumna("gross_revenue_pln"))
    koszt = sum(Decimal(str(value)) for value in kolumna("cost_pln")) + sum(
        Decimal(str(value)) for value in kolumna("sell_commission_alloc_pln")
    )
    podsumowanie = result.annual_summary

    assert przychod == Decimal(str(podsumowanie["total_revenue_pln"]))
    assert koszt == Decimal(str(podsumowanie["total_cost_pln"]))

    # Wiersze innych lat nadal sa w pakiecie - potrzebne przy stratach - ale
    # jawnie odseparowane, zeby nie wchodzily do sumy uzgadnianej z formularzem.
    inne = workbook["FIFO_Other_Years"]
    lata_innych = {row[naglowki.index("tax_year")] for row in inne.iter_rows(min_row=2, values_only=True)}
    assert rok not in lata_innych
